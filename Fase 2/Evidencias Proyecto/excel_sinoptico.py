"""Exportación de un sinóptico a Excel (.xlsx) con openpyxl."""
import io
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

DIAS = [(1, "Lunes"), (2, "Martes"), (3, "Miércoles"), (4, "Jueves"), (5, "Viernes"), (6, "Sábado")]
NAVY = "1F2A44"
BORDE = Border(*(Side(style="thin", color="BFC5D0"),) * 4)


def _hhmm(t):
    return t.strftime("%H:%M") if hasattr(t, "strftime") else str(t)[:5]


def _texto_semanas(semanas):
    if not semanas:
        return ""
    if len(semanas) == 17:
        return "Todas (1-17)"
    return ", ".join(str(x) for x in semanas)


def _color_texto(hex_color):
    """Blanco sobre colores oscuros, negro sobre claros."""
    try:
        h = hex_color.lstrip("#")
        r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
        return "FFFFFF" if (0.299 * r + 0.587 * g + 0.114 * b) < 150 else "000000"
    except Exception:
        return "000000"


def construir_excel_sinoptico(*, sinoptico_id, carrera_nombre, jornada_nombre, semestre_numero, anio, items, bloques_horas):
    wb = Workbook()

    # ---------------- Hoja 1: grilla ----------------
    ws = wb.active
    ws.title = "Horario"
    titulo = f"Sinóptico N° {sinoptico_id} — {carrera_nombre} — Jornada {jornada_nombre} — Semestre {semestre_numero}/{anio}"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(DIAS) + 1)
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=13, color=NAVY)
    ws.row_dimensions[1].height = 24

    ws.cell(row=3, column=1, value="Módulo").font = Font(bold=True, color="FFFFFF")
    ws.cell(row=3, column=1).fill = PatternFill("solid", fgColor=NAVY)
    for c, (_, nombre) in enumerate(DIAS, start=2):
        cel = ws.cell(row=3, column=c, value=nombre)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor=NAVY)
        cel.alignment = Alignment(horizontal="center", vertical="center")
        cel.border = BORDE

    franjas = sorted({(a, b) for a, b in bloques_horas})
    por_celda = {}
    for it in items:
        por_celda.setdefault((it["dia_semana"], it["hora_inicio"], it["hora_fin"]), []).append(it)

    for r, (ini, fin) in enumerate(franjas, start=4):
        cel = ws.cell(row=r, column=1, value=f"{_hhmm(ini)} - {_hhmm(fin)}")
        cel.font = Font(bold=True)
        cel.alignment = Alignment(horizontal="center", vertical="center")
        cel.border = BORDE
        ws.row_dimensions[r].height = 48
        for c, (dia, _) in enumerate(DIAS, start=2):
            cel = ws.cell(row=r, column=c)
            cel.border = BORDE
            lista = por_celda.get((dia, ini, fin), [])
            if not lista:
                continue
            cel.value = "\n\n".join(
                f"{x['nombre']}" + (f" ({x['tipo']})" if x["tipo"] else "")
                + f"\n{x['profesor'] or 'Sin profesor'}\n{x['sala'] or 'Sin sala'}"
                for x in lista
            )
            color = (lista[0].get("color") or "#2563EB").lstrip("#").upper()
            cel.fill = PatternFill("solid", fgColor=color)
            cel.font = Font(size=9, color=_color_texto(color))
            cel.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    ws.column_dimensions["A"].width = 16
    for c in range(2, len(DIAS) + 2):
        ws.column_dimensions[ws.cell(row=3, column=c).column_letter].width = 28
    ws.freeze_panes = "B4"

    # ---------------- Hoja 2: detalle ----------------
    wd = wb.create_sheet("Detalle")
    cab = ["Sección", "Código", "Asignatura", "Tipo", "Profesor", "Sala", "Capacidad", "Día", "Hora inicio", "Hora fin", "Semanas"]
    for c, t in enumerate(cab, start=1):
        cel = wd.cell(row=1, column=c, value=t)
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor=NAVY)
        cel.alignment = Alignment(horizontal="center")
    nombres_dia = dict(DIAS)
    orden = sorted(items, key=lambda x: (x["dia_semana"], x["hora_inicio"], x["nombre"]))
    for r, x in enumerate(orden, start=2):
        fila = [
            x["seccion"] or "", x["codigo"], x["nombre"], x["tipo"] or "", x["profesor"] or "", x["sala"] or "",
            x["capacidad"] if x["capacidad"] is not None else "",
            nombres_dia.get(x["dia_semana"], x["dia_semana"]), _hhmm(x["hora_inicio"]), _hhmm(x["hora_fin"]),
            _texto_semanas(x["semanas"]),
        ]
        for c, v in enumerate(fila, start=1):
            wd.cell(row=r, column=c, value=v)
    for letra, ancho in zip("ABCDEFGHIJK", [14, 12, 40, 10, 28, 12, 11, 12, 12, 10, 24]):
        wd.column_dimensions[letra].width = ancho
    wd.freeze_panes = "A2"
    wd.auto_filter.ref = f"A1:K{max(len(orden) + 1, 2)}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()