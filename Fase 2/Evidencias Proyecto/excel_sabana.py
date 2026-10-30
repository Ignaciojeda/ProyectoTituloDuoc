from collections import defaultdict
import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

DIAS = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado"]
GAP_MAX_MIN = 5  # entre módulos pegados hay ~1 min; el recreo es de ~10

COLUMNAS = [
    "Año académico", "Período académico", "Sede", "Nombre Sede",
    "Clase de plan de estudios", "Plan Estudio", "Escuela",
    "Cod.Jornada", "Nom.Jornada", "Nivel Semestral",
    "ID seccion", "Seccion", "Nombre Seccion",
    "Cod Asignatura", "Nom. Asignatura", "Formacion",
    "Rut Docente", "Instructor de un evento (denominación)",
    "Denominación de sala", "Capacidad Sala",
    "Hora de inicio", "Hora final",
    *DIAS,
    "Duración horas", "Fecha de inicio", "Fecha final",
    "Fecha inicio elemento programa", "Fecha final elemento programa",
    "Capacidad Inicial",
]
ANCHAS = {"Escuela", "Nombre Seccion", "Nom. Asignatura",
          "Instructor de un evento (denominación)", "Denominación de sala"}


def _minutos(t):
    return t.hour * 60 + t.minute


def _fusionar(items):
    """Une los módulos pegados (sin recreo) de una misma clase en un solo tramo."""
    grupos = defaultdict(list)
    for it in items:
        k = (it["sinoptico_id"], it["asignatura_id"], it["profesor_id"], it["sala_id"], it["dia_semana"])
        grupos[k].append(it)
    filas = []
    for lista in grupos.values():
        lista.sort(key=lambda x: x["hora_inicio"])
        actual = dict(lista[0])
        for it in lista[1:]:
            if _minutos(it["hora_inicio"]) - _minutos(actual["hora_fin"]) <= GAP_MAX_MIN:
                actual["hora_fin"] = it["hora_fin"]
            else:
                filas.append(actual)
                actual = dict(it)
        filas.append(actual)
    return filas


def construir_excel_sabana(items):
    # fechas de la sección = rango de todas sus clases
    rango = {}
    for it in items:
        if not it["fecha_inicio"] or not it["fecha_final"]:
            continue
        k = (it["sinoptico_id"], it["asignatura_id"])
        if k in rango:
            rango[k] = (min(rango[k][0], it["fecha_inicio"]), max(rango[k][1], it["fecha_final"]))
        else:
            rango[k] = (it["fecha_inicio"], it["fecha_final"])

    por_semestre = defaultdict(list)
    for it in items:
        por_semestre[(it["anio"], it["periodo"])].append(it)

    wb = Workbook()
    wb.remove(wb.active)
    if not por_semestre:
        wb.create_sheet("Sábana")

    for (anio, periodo), lista in sorted(por_semestre.items()):
        ws = wb.create_sheet(f"{anio}_{periodo}")
        ws.append(COLUMNAS)

        filas = _fusionar(lista)
        filas.sort(key=lambda x: (x["escuela"], x["nivel"] or 0, x["seccion"] or "", x["dia_semana"], x["hora_inicio"]))

        for f in filas:
            ini, fin = rango.get((f["sinoptico_id"], f["asignatura_id"]), (None, None))
            ws.append([
                f["anio"], f["periodo"], "52000101", "PUERTO MONTT",
                f["cod_carrera"], f["plan"], f["escuela"],
                f["cod_jornada"], f["nom_jornada"], f["nivel"],
                f["id_seccion"], f["seccion"], f["nombre_seccion"],
                f["cod_asignatura"], f["nom_asignatura"], f["formacion"],
                f["rut"], f["profesor"],
                f["sala"], f["capacidad_sala"],
                f["hora_inicio"], f["hora_fin"],
                *["X" if f["dia_semana"] == i else "" for i in range(1, 7)],
                f["duracion_horas"], ini, fin,
                f["fecha_inicio"], f["fecha_final"],
                f["capacidad_inicial"],
            ])

        # formato
        cab = {c: i for i, c in enumerate(COLUMNAS)}
        for row in ws.iter_rows(min_row=2):
            for n in ("Hora de inicio", "Hora final"):
                row[cab[n]].number_format = "hh:mm"
            for n in ("Fecha de inicio", "Fecha final", "Fecha inicio elemento programa", "Fecha final elemento programa"):
                row[cab[n]].number_format = "dd-mm-yyyy"
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3864")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for i, n in enumerate(COLUMNAS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = 38 if n in ANCHAS else 15
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()